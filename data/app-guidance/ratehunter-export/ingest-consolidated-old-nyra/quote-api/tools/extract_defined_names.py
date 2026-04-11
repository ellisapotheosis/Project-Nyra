"""Extract defined names (named ranges + named formulas) from your Excel mortgage calculator.

Why this exists:
- Your original workbook is .xls; converting to .xlsx can mangle long formulas.
- But named ranges / named formulas still give us a strong "semantic map" of what the sheet expects.

Usage:
  python tools/extract_defined_names.py \
    --input "/path/to/All in One Calculator.xls" \
    --output defined_names.json

Requires:
- LibreOffice (`soffice`) in PATH (present in our docker stack images by default)
- openpyxl
"""

from __future__ import annotations

import argparse
import json
import os
import subprocess
import tempfile

from openpyxl import load_workbook


def convert_xls_to_xlsx(xls_path: str) -> str:
    outdir = tempfile.mkdtemp(prefix="nyra_xls_")
    subprocess.check_call(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            outdir,
            xls_path,
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    base = os.path.splitext(os.path.basename(xls_path))[0]
    xlsx_path = os.path.join(outdir, base + ".xlsx")
    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(f"Conversion failed; expected {xlsx_path}")
    return xlsx_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args()

    inp = args.input
    if inp.lower().endswith(".xls") and not inp.lower().endswith(".xlsx"):
        xlsx = convert_xls_to_xlsx(inp)
    else:
        xlsx = inp

    wb = load_workbook(xlsx, data_only=False)
    out = {
        "sheetnames": wb.sheetnames,
        "defined_names": {},
    }

    for name in wb.defined_names.keys():
        dn = wb.defined_names[name]
        # dn.attr_text is either a range like 'Sheet!$A$1' or a formula.
        out["defined_names"][name] = dn.attr_text

    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2, ensure_ascii=False)

    print(f"Wrote {args.output} ({len(out['defined_names'])} defined names)")


if __name__ == "__main__":
    main()
